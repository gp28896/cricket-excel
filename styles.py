"""
styles.py
=========

Reusable NamedStyle registration and styling helpers for the Cricket
Tournament Excel system.

Overview
--------
This module is responsible for defining and registering the reusable
``openpyxl`` styling building blocks — fonts, fills, borders, alignments,
and ``NamedStyle`` objects — used across the workbook.

Unlike higher-level modules such as ``workbook.py``, ``create_match.py``,
or ``scorecard.py``, which orchestrate worksheet construction and business
logic, this module focuses solely on the visual presentation layer. It
provides the shared styling vocabulary that other modules apply to cells,
ranges, and worksheets rather than each module defining its own ad-hoc
formatting.

Responsibilities
-----------------
The responsibilities of this module include (implemented incrementally
across future parts of this file):

* NamedStyle definition and registration
* Font, fill, border, and alignment presets
* Header, title, and label styling helpers
* Conditional formatting style presets
* Theme/color palette constants
* Style application helpers for cells and ranges
* Protection style presets
* Number format presets

Compatibility
--------------
This module has been designed to integrate with:

* ``constants.py`` -- application constants and configuration values
* ``workbook.py`` -- ``WorkbookBuilder._register_named_style`` looks up
  factory functions in this module by name (e.g.
  ``create_header_named_style``) and calls them with the active
  ``WorkbookConfig`` instance, falling back to a no-argument call if the
  factory does not accept one
* ``create_match.py`` -- consumes styling helpers via ``from styles
  import *``
* ``scorecard.py`` -- consumes styling helpers for scorecard formatting

This module has no dependency on ``formulas.py`` or ``validation.py``
and does not perform any workbook orchestration, business logic, or data
validation itself. It also has no hard dependency on ``workbook.py`` --
configuration objects are consumed structurally (duck-typed) rather than
through a direct import, to avoid circular imports.

Design Philosophy
------------------
The architecture intentionally separates styling from orchestration.

Every visual style used in the workbook should be defined once, here, and
reused everywhere it is needed rather than being redefined inline inside
worksheet-building code. Each named-style factory:

* accepts an optional configuration object and falls back to sensible
  defaults when one is not supplied or does not expose an expected
  attribute
* returns a fully-formed :class:`~openpyxl.styles.NamedStyle` ready to be
  registered onto a workbook
* has no side effects and does not mutate any workbook or worksheet
  directly -- registration remains the responsibility of ``workbook.py``

This design provides:

* a single source of truth for visual presentation
* consistent look and feel across all worksheets
* easier theme-wide changes
* simpler testing of style definitions in isolation
* reduced duplication across modules
"""

from __future__ import annotations

from typing import Protocol
from typing import runtime_checkable

###############################################################################
# Third-party imports
###############################################################################

from openpyxl.styles import Alignment
from openpyxl.styles import Border
from openpyxl.styles import Font
from openpyxl.styles import NamedStyle
from openpyxl.styles import PatternFill
from openpyxl.styles import Protection
from openpyxl.styles import Side

###############################################################################
# Configuration protocol
###############################################################################


@runtime_checkable
class _WorkbookConfigLike(Protocol):
    """Structural type describing the subset of ``WorkbookConfig`` used here.

    Factories accept any object exposing these attributes (typically a
    ``workbook.WorkbookConfig`` instance) without requiring a hard import
    of ``workbook.py``, keeping this module dependency-free.
    """

    default_font: str
    default_font_size: float
    date_format: str
    time_format: str


###############################################################################
# Theme defaults
###############################################################################

_DEFAULT_FONT_FAMILY = "Calibri"
_DEFAULT_FONT_SIZE = 11.0
_DEFAULT_DATE_FORMAT = "yyyy-mm-dd"
_DEFAULT_TIME_FORMAT = "hh:mm:ss"

_HEADER_FILL_COLOR = "D9D9D9"
_BORDER_COLOR = "D9D9D9"
_WHITE = "FFFFFF"
_BLACK = "000000"


def _font_family(config: _WorkbookConfigLike | None) -> str:
    """Resolve the font family to use, falling back to the theme default."""
    return getattr(config, "default_font", None) or _DEFAULT_FONT_FAMILY


def _font_size(config: _WorkbookConfigLike | None) -> float:
    """Resolve the font size to use, falling back to the theme default."""
    return getattr(config, "default_font_size", None) or _DEFAULT_FONT_SIZE


def _date_format(config: _WorkbookConfigLike | None) -> str:
    """Resolve the date number format, falling back to the theme default."""
    return getattr(config, "date_format", None) or _DEFAULT_DATE_FORMAT


def _time_format(config: _WorkbookConfigLike | None) -> str:
    """Resolve the time number format, falling back to the theme default."""
    return getattr(config, "time_format", None) or _DEFAULT_TIME_FORMAT


def _standard_border() -> Border:
    """Create the standard thin-bordered box used across preset styles."""
    side = Side(style="thin", color=_BORDER_COLOR)
    return Border(left=side, right=side, top=side, bottom=side)


###############################################################################
# Named style factories
###############################################################################
#
# Each factory is looked up by name via `getattr` inside
# `WorkbookBuilder._register_named_style` (see workbook.py) and invoked as
# `factory(config)`, falling back to `factory()` if the factory does not
# accept an argument. Every factory below accepts an optional config and
# is safe to call either way.


def create_header_named_style(
    config: _WorkbookConfigLike | None = None,
) -> NamedStyle:
    """Create the standard worksheet header style.

    Args:
        config:
            Optional workbook configuration used to resolve font family and
            size. Falls back to theme defaults when omitted.

    Returns:
        A fully-formed ``NamedStyle`` named ``"Header"``.
    """
    style = NamedStyle(name="Header")
    style.font = Font(
        name=_font_family(config),
        size=_font_size(config),
        bold=True,
        color=_BLACK,
    )
    style.fill = PatternFill(
        fill_type="solid",
        fgColor=_HEADER_FILL_COLOR,
        bgColor=_HEADER_FILL_COLOR,
    )
    style.border = _standard_border()
    style.alignment = Alignment(
        horizontal="center",
        vertical="center",
        wrap_text=True,
        shrink_to_fit=False,
    )
    style.number_format = "General"
    return style


def create_title_named_style(
    config: _WorkbookConfigLike | None = None,
) -> NamedStyle:
    """Create the workbook/worksheet title style.

    Args:
        config:
            Optional workbook configuration used to resolve font family and
            size. Falls back to theme defaults when omitted.

    Returns:
        A fully-formed ``NamedStyle`` named ``"Title"``.
    """
    style = NamedStyle(name="Title")
    style.font = Font(
        name=_font_family(config),
        size=_font_size(config) + 6,
        bold=True,
        color=_BLACK,
    )
    style.fill = PatternFill(fill_type="solid", fgColor=_WHITE, bgColor=_WHITE)
    style.border = Border()
    style.alignment = Alignment(
        horizontal="center",
        vertical="center",
        wrap_text=False,
        shrink_to_fit=False,
    )
    style.number_format = "General"
    return style


def create_normal_named_style(
    config: _WorkbookConfigLike | None = None,
) -> NamedStyle:
    """Create the default workbook style applied to ordinary cells.

    Args:
        config:
            Optional workbook configuration used to resolve font family and
            size. Falls back to theme defaults when omitted.

    Returns:
        A fully-formed ``NamedStyle`` named ``"Normal"``.
    """
    style = NamedStyle(name="Normal")
    style.font = Font(
        name=_font_family(config),
        size=_font_size(config),
        bold=False,
        color=_BLACK,
    )
    style.fill = PatternFill(fill_type="solid", fgColor=_WHITE, bgColor=_WHITE)
    style.border = Border()
    style.alignment = Alignment(
        horizontal="general",
        vertical="center",
        wrap_text=False,
        shrink_to_fit=False,
    )
    style.number_format = "General"
    return style


def create_number_named_style(
    config: _WorkbookConfigLike | None = None,
) -> NamedStyle:
    """Create the standard numeric style.

    Args:
        config:
            Optional workbook configuration used to resolve font family and
            size. Falls back to theme defaults when omitted.

    Returns:
        A fully-formed ``NamedStyle`` named ``"Number"``.
    """
    style = NamedStyle(name="Number")
    style.font = Font(
        name=_font_family(config),
        size=_font_size(config),
        bold=False,
        color=_BLACK,
    )
    style.fill = PatternFill(fill_type="solid", fgColor=_WHITE, bgColor=_WHITE)
    style.border = Border()
    style.alignment = Alignment(
        horizontal="right",
        vertical="center",
        wrap_text=False,
        shrink_to_fit=False,
    )
    style.number_format = "0"
    return style


def create_percentage_named_style(
    config: _WorkbookConfigLike | None = None,
) -> NamedStyle:
    """Create the standard percentage style.

    Args:
        config:
            Optional workbook configuration used to resolve font family and
            size. Falls back to theme defaults when omitted.

    Returns:
        A fully-formed ``NamedStyle`` named ``"Percentage"``.
    """
    style = NamedStyle(name="Percentage")
    style.font = Font(
        name=_font_family(config),
        size=_font_size(config),
        bold=False,
        color=_BLACK,
    )
    style.fill = PatternFill(fill_type="solid", fgColor=_WHITE, bgColor=_WHITE)
    style.border = Border()
    style.alignment = Alignment(
        horizontal="right",
        vertical="center",
        wrap_text=False,
        shrink_to_fit=False,
    )
    style.number_format = "0.00%"
    return style


def create_date_named_style(
    config: _WorkbookConfigLike | None = None,
) -> NamedStyle:
    """Create the standard date style.

    Args:
        config:
            Optional workbook configuration used to resolve the date number
            format and font. Falls back to theme defaults when omitted.

    Returns:
        A fully-formed ``NamedStyle`` named ``"Date"``.
    """
    style = NamedStyle(name="Date")
    style.font = Font(
        name=_font_family(config),
        size=_font_size(config),
        bold=False,
        color=_BLACK,
    )
    style.fill = PatternFill(fill_type="solid", fgColor=_WHITE, bgColor=_WHITE)
    style.border = Border()
    style.alignment = Alignment(
        horizontal="center",
        vertical="center",
        wrap_text=False,
        shrink_to_fit=False,
    )
    style.number_format = _date_format(config)
    return style


def create_time_named_style(
    config: _WorkbookConfigLike | None = None,
) -> NamedStyle:
    """Create the standard time style.

    Args:
        config:
            Optional workbook configuration used to resolve the time number
            format and font. Falls back to theme defaults when omitted.

    Returns:
        A fully-formed ``NamedStyle`` named ``"Time"``.
    """
    style = NamedStyle(name="Time")
    style.font = Font(
        name=_font_family(config),
        size=_font_size(config),
        bold=False,
        color=_BLACK,
    )
    style.fill = PatternFill(fill_type="solid", fgColor=_WHITE, bgColor=_WHITE)
    style.border = Border()
    style.alignment = Alignment(
        horizontal="center",
        vertical="center",
        wrap_text=False,
        shrink_to_fit=False,
    )
    style.number_format = _time_format(config)
    return style


def create_currency_named_style(
    config: _WorkbookConfigLike | None = None,
) -> NamedStyle:
    """Create the standard currency style.

    Args:
        config:
            Optional workbook configuration used to resolve font family and
            size. Falls back to theme defaults when omitted.

    Returns:
        A fully-formed ``NamedStyle`` named ``"Currency"``.
    """
    style = NamedStyle(name="Currency")
    style.font = Font(
        name=_font_family(config),
        size=_font_size(config),
        bold=False,
        color=_BLACK,
    )
    style.fill = PatternFill(fill_type="solid", fgColor=_WHITE, bgColor=_WHITE)
    style.border = Border()
    style.alignment = Alignment(
        horizontal="right",
        vertical="center",
        wrap_text=False,
        shrink_to_fit=False,
    )
    style.number_format = "#,##0.00"
    return style


def create_hidden_named_style(
    config: _WorkbookConfigLike | None = None,
) -> NamedStyle:
    """Create the style used for hidden/internal worksheet cells.

    The font color matches the fill color so content remains present but
    visually invisible, and the cell is locked via protection.

    Args:
        config:
            Optional workbook configuration used to resolve font family and
            size. Falls back to theme defaults when omitted.

    Returns:
        A fully-formed ``NamedStyle`` named ``"Hidden"``.
    """
    style = NamedStyle(name="Hidden")
    style.font = Font(
        name=_font_family(config),
        size=_font_size(config),
        bold=False,
        color=_WHITE,
    )
    style.fill = PatternFill(fill_type="solid", fgColor=_WHITE, bgColor=_WHITE)
    style.border = Border()
    style.alignment = Alignment(
        horizontal="general",
        vertical="center",
        wrap_text=False,
        shrink_to_fit=False,
    )
    style.protection = Protection(locked=True, hidden=True)
    style.number_format = "General"
    return style


###############################################################################
# Public API
###############################################################################

__all__: list[str] = [
    "create_header_named_style",
    "create_title_named_style",
    "create_normal_named_style",
    "create_number_named_style",
    "create_percentage_named_style",
    "create_date_named_style",
    "create_time_named_style",
    "create_currency_named_style",
    "create_hidden_named_style",
]